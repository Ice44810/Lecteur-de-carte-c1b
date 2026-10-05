"""Export Excel d'un rapport : une feuille par rubrique."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.reports.tables import ReportTable

__all__ = ["write_excel"]

_TITLE_FONT = Font(bold=True, size=13)
_HEADER_FONT = Font(bold=True)
_EMPHASIS_FILL = PatternFill("solid", fgColor="BFE8E8")
_SHEET_NAME_MAX = 31


def write_excel(path: Path, header: str, tables: tuple[ReportTable, ...]) -> Path:
    """Ecrit un classeur Excel.

    Args:
        path: Fichier de destination (``.xlsx``).
        header: En-tete repris en premiere ligne de chaque feuille.
        tables: Rubriques du rapport, une feuille chacune.

    Returns:
        Le chemin du fichier ecrit.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)
    for table in tables:
        sheet = workbook.create_sheet(table.title[:_SHEET_NAME_MAX])
        sheet.append([header])
        sheet["A1"].font = _HEADER_FONT
        sheet.append([table.title])
        sheet["A2"].font = _TITLE_FONT
        sheet.append([])
        sheet.append(list(table.headers))
        for cell in sheet[4]:
            cell.font = _HEADER_FONT
        for index, row in enumerate(table.rows):
            sheet.append(list(row))
            if index in table.emphasized:
                for cell in sheet[sheet.max_row]:
                    cell.fill = _EMPHASIS_FILL
        if table.note:
            sheet.append([])
            sheet.append([table.note])
        sheet.freeze_panes = "A5"
        for column, title in enumerate(table.headers, start=1):
            values = [title, *(row[column - 1] for row in table.rows if column <= len(row))]
            width = min(max(len(str(value)) for value in values) + 2, 60)
            sheet.column_dimensions[get_column_letter(column)].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
